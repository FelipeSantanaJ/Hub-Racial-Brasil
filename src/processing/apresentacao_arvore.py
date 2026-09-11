"""Monta o deck da ÁRVORE cruzada de renda por raça a partir de
`docs/arvore_manifest.json`.

Cada folha da árvore vira 1 slide: a IMAGEM é "crua" (só dados/eixos/rótulos de
valor — ver `graficos_arvore.py`) e o TEXTO (título, subtítulo, legenda, linha de
fonte) é caixa de texto editável no PowerPoint, não rasterizado.

Enquanto Fase D (topo10%/quartis) e Fase F (deck final consolidado com "Outras
análises") não fecham, este arquivo (`Datahub_Racial_Brasil_ARVORE.pptx`) é o
deck de conferência da estrutura nova.

Uso:
    python -m src.processing.apresentacao_arvore
"""
from __future__ import annotations

import json

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image

from src.processing.graficos_fase1 import REPO_ROOT

IMG_DIR = REPO_ROOT / "docs" / "img"
MANIFESTO = REPO_ROOT / "docs" / "arvore_manifest.json"
DESTINO = REPO_ROOT / "docs" / "Datahub_Racial_Brasil_ARVORE.pptx"

LARGURA, ALTURA = Inches(13.333), Inches(7.5)
SUPERFICIE = RGBColor(0xF7, 0xF2, 0xEA)
TINTA_PRIMARIA = RGBColor(0x2B, 0x20, 0x18)
TINTA_SECUNDARIA = RGBColor(0x5C, 0x4F, 0x3F)
TINTA_MUTED = RGBColor(0x8F, 0x82, 0x71)
ACENTO = RGBColor(0x0D, 0x90, 0x86)

FONTE_PADRAO = "Fonte: IBGE, PNAD Contínua Trimestral (microdados). Elaboração própria."


def _slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = SUPERFICIE
    return s


def _txt(slide, left, top, width, height, texto, tamanho, cor, *, bold=False,
         align=PP_ALIGN.LEFT, italic=False):
    cx = slide.shapes.add_textbox(left, top, width, height)
    tf = cx.text_frame
    tf.word_wrap = True
    for i, linha in enumerate(texto.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = linha
        p.font.size = Pt(tamanho)
        p.font.bold = bold
        p.font.italic = italic
        p.font.color.rgb = cor
        p.alignment = align
    return cx


def slide_titulo(prs):
    s = _slide(prs)
    _txt(s, Inches(0.8), Inches(2.5), Inches(11.7), Inches(1.2),
         "Datahub de Análise Racial no Brasil", 40, TINTA_PRIMARIA, bold=True)
    _txt(s, Inches(0.8), Inches(3.7), Inches(11.7), Inches(0.9),
         "Renda por raça — árvore cruzada (Raça × Gênero × Faixa/Geração × Escolaridade)",
         20, TINTA_SECUNDARIA)
    _txt(s, Inches(0.8), Inches(4.5), Inches(11.7), Inches(0.6),
         "PNAD Contínua Trimestral (2012–2026) · cada gráfico aparece em um único slide",
         13, TINTA_MUTED)


def slide_indice(prs, secoes):
    s = _slide(prs)
    _txt(s, Inches(0.6), Inches(0.4), Inches(12.1), Inches(0.7),
         "Índice — árvore de renda por raça", 22, ACENTO, bold=True)
    linhas = [f"{sec['item']}. {sec['titulo']}  —  {len(sec['folhas'])} slides" for sec in secoes]
    meio = (len(linhas) + 1) // 2
    for i, col in enumerate((linhas[:meio], linhas[meio:])):
        _txt(s, Inches(0.6) + i * Inches(6.4), Inches(1.3), Inches(6.1), Inches(5.6),
             "\n".join(col), 13, TINTA_PRIMARIA)


def slide_divisor(prs, titulo, nota):
    s = _slide(prs)
    _txt(s, Inches(0.8), Inches(3.0), Inches(11.7), Inches(1.2), titulo, 32, ACENTO, bold=True)
    if nota:
        _txt(s, Inches(0.8), Inches(4.2), Inches(11.7), Inches(1.0), nota, 15, TINTA_SECUNDARIA)


def slide_subdivisor(prs, titulo):
    s = _slide(prs)
    _txt(s, Inches(0.9), Inches(3.3), Inches(11.5), Inches(0.9), titulo, 24, TINTA_SECUNDARIA, bold=True)


def ramo_branch(ramo: str) -> str | None:
    """Só 'Mulheres'/'Homens' (itens 6-10) valem um slide subdivisor. 'Combinado' e '—'
    são folhas dentro do item, não uma ramificação de verdade."""
    if not ramo or ramo == "—":
        return None
    prim = ramo.split(" · ")[0].split(" ")[0]
    return prim if prim in ("Mulheres", "Homens") else None


def slide_grafico(prs, folha):
    s = _slide(prs)
    _txt(s, Inches(0.55), Inches(0.32), Inches(12.2), Inches(0.7),
         folha["titulo"], 21, TINTA_PRIMARIA, bold=True)
    _txt(s, Inches(0.55), Inches(1.02), Inches(12.2), Inches(0.55),
         folha["subtitulo"], 11.5, TINTA_SECUNDARIA)

    caminho = IMG_DIR / folha["arquivo"]
    if caminho.exists():
        with Image.open(caminho) as im:
            prop = im.size[0] / im.size[1]
        max_w, max_h = Inches(12.0), Inches(4.95)
        if max_w / prop <= max_h:
            w, h = max_w, Emu(int(max_w / prop))
        else:
            h, w = max_h, Emu(int(max_h * prop))
        s.shapes.add_picture(str(caminho), Emu(int((LARGURA - w) / 2)), Inches(1.68), width=w, height=h)
    else:
        _txt(s, Inches(0.55), Inches(3.5), Inches(12.2), Inches(0.6),
             f"[imagem ausente: {folha['arquivo']}]", 12, TINTA_MUTED)

    _txt(s, Inches(0.7), Inches(6.72), Inches(11.9), Inches(0.55),
         folha["legenda"], 11.5, TINTA_SECUNDARIA, align=PP_ALIGN.CENTER, italic=True)
    _txt(s, Inches(0.55), Inches(7.16), Inches(9.0), Inches(0.3),
         FONTE_PADRAO, 8, TINTA_MUTED)


NOTA_ITEM = {
    1: "Renda por raça (Branca / Negra=Preta+Parda / Indígena): série 2012–2026. Hiato: série separada por comparação (vs. Negra / vs. Indígena); o snapshot é 1 gráfico com as 2 barras. Série de Indígena suavizada (amostra pequena).",
    2: "Renda: série combinada (Homens tracejado, Mulheres sólido) + Mulheres + Homens. Hiato: série (combinada + por gênero) e snapshot, cada comparação em seu slide.",
    3: "Renda: série combinada em 3 slides (um por raça) + uma por faixa. Hiato: série (combinada + por faixa, com a faixa no título) e snapshot, cada comparação em seu slide.",
    4: "Idem por geração (mesma coorte de nascimento).",
    5: "Idem por nível de instrução (7 níveis).",
    6: "Quebra CRUZADA gênero×faixa. Por gênero: renda (série combinada + uma por faixa) + hiato por comparação (série combinada, série por faixa, snapshot).",
    7: "Quebra CRUZADA gênero×geração.",
    8: "Quebra CRUZADA gênero×escolaridade.",
    9: "Quebra CRUZADA gênero×faixa×escolaridade — 70 células (série de renda + série de hiato Branca-vs-Negra por célula). Indígena via 2 heatmaps por gênero (vs. Negra / vs. Indígena).",
    10: "Quebra CRUZADA gênero×geração×escolaridade — 56 células.",
    11: "Geração × Escolaridade com AMBOS os gêneros juntos — 28 células (série de renda + série de hiato por célula) + 2 heatmaps. Fecha a combinação que faltava pra um dashboard geração×escolaridade×gênero.",
}


def main():
    manifesto = json.loads(MANIFESTO.read_text(encoding="utf-8"))
    secoes = manifesto["arvore"]

    prs = Presentation()
    prs.slide_width, prs.slide_height = LARGURA, ALTURA
    slide_titulo(prs)
    slide_indice(prs, secoes)

    n = 2
    for sec in secoes:
        n_val = sum(1 for f in sec["folhas"] if f["tipo"].startswith("valor"))
        n_hi = sum(1 for f in sec["folhas"] if f["tipo"].startswith("hiato"))
        slide_divisor(prs, f"{sec['item']}. {sec['titulo']}",
                      f"{n_val} de renda · {n_hi} de hiato · " + NOTA_ITEM.get(sec["item"], ""))
        n += 1
        ramo_atual = None
        for folha in sec["folhas"]:
            br = ramo_branch(folha["ramo"])
            if br and br != ramo_atual:
                slide_subdivisor(prs, f"{sec['titulo']} — {br}")
                n += 1
            ramo_atual = br
            slide_grafico(prs, folha)
            n += 1

    prs.save(DESTINO)
    print(f"Deck salvo: {DESTINO.relative_to(REPO_ROOT)}")
    print(f"Slides: {n}  ({sum(len(s['folhas']) for s in secoes)} folhas + divisores + título/índice)")


if __name__ == "__main__":
    main()
