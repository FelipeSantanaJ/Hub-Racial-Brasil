"""Gera a apresentação em PowerPoint da Fase 1, com todos os gráficos de
docs/img/ organizados por seção (mesma estrutura de docs/ANALISE_FASE1.md).

Uso:
    python -m src.processing.apresentacao_fase1
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

REPO_ROOT = Path(__file__).resolve().parents[2]
IMG_DIR = REPO_ROOT / "docs" / "img"
DESTINO = REPO_ROOT / "docs" / "Datahub_Racial_Brasil_Fase1.pptx"

LARGURA_SLIDE = Inches(13.333)
ALTURA_SLIDE = Inches(7.5)


# Identidade visual em tons terrosos — mesma paleta de src/processing/graficos_fase1.py.
SUPERFICIE = RGBColor(0xF7, 0xF2, 0xEA)
TINTA_PRIMARIA = RGBColor(0x2B, 0x20, 0x18)
TINTA_SECUNDARIA = RGBColor(0x5C, 0x4F, 0x3F)
COR_BRANCA = RGBColor(0x0D, 0x90, 0x86)

# (título da seção, nota de rodapé opcional, [(arquivo, legenda), ...])
SECOES = [
    ("Raça", None, [
        ("renda_por_raca.png", "Hiato grande e persistente entre Branca e Negra/Indígena; estreitamento leve 2012→2026."),
        ("renda_preta_parda.png", "Preta e Parda andam próximas, com Parda ultrapassando levemente nos últimos anos."),
    ]),
    ("Hiato Branca vs. Negra — série histórica", "Com teste de significância (Welch, IC 95%)", [
        ("hiato_racial_percentual.png", "Hiato relativo caiu de ~76% para ~66% — todos os 58 trimestres são estatisticamente significativos."),
        ("hiato_racial_absoluto.png", "Em R$, o hiato disparou na pandemia (pico ~R$2.035) antes de recuar."),
    ]),
    ("É ocupação, ou é cor da pele?", "Decomposição: quanto idade, escolaridade e ocupação explicam do hiato", [
        ("decomposicao_hiato_ocupacional.png", "Hiato bruto 67% → 64% (só idade) → 30% (+escolaridade) → 24% (+ocupação, residual). Ocupação SOZINHA (isolada, sem outros controles) já leva o hiato a 33%."),
        ("oaxaca_blinder_decomposicao.png", "Mesma pergunta via regressão (Oaxaca-Blinder), com teste de significância: resíduo de +22% (p<0,001) — bate com a padronização direta. Ocupação sozinha explica 40% do hiato (p<0,001)."),
        ("oaxaca_blinder_quantis.png", "O hiato residual não é uniforme: menor na mediana (~15%), maior na base (~31%, 'piso pegajoso') e no topo (~36%, 'teto de vidro')."),
        ("renda_por_raca_ocupacao.png", "Dentro de CADA categoria ocupacional, Branca ganha mais que Negra — o hiato não é só 'estar em ocupações diferentes'."),
    ]),
    ("Raça cruzada com tudo: as combinações que faltavam", "Heatmaps raça × 2 dimensões, um painel por raça — fecha a matriz completa pedida pelo usuário", [
        ("raca_genero_geracao.png", "Renda por raça, gênero e geração."),
        ("raca_genero_ocupacao.png", "Renda por raça, gênero e ocupação."),
        ("raca_faixa_etaria_escolaridade.png", "Renda por raça, faixa etária e escolaridade."),
        ("raca_faixa_etaria_ocupacao.png", "Renda por raça, faixa etária e ocupação."),
        ("raca_geracao_escolaridade.png", "Renda por raça, geração e escolaridade."),
        ("raca_geracao_ocupacao.png", "Renda por raça, geração e ocupação."),
        ("raca_escolaridade_ocupacao.png", "Renda por raça, escolaridade e ocupação — dentro do mesmo nível de instrução E mesma ocupação, o hiato racial ainda aparece na maioria das células."),
    ]),
    ("Aprofundamentos: região, segregação e quebras estruturais", "Rodada 3 — a pedido do usuário", [
        ("hiato_racial_por_regiao.png", "O hiato racial varia MUITO por região: ~65-85% no Sudeste vs. ~45-55% no Sul."),
        ("segregacao_ocupacional.png", "Índice de Duncan: ~17% da população Branca ou Negra precisaria trocar de categoria ocupacional pra igualar a distribuição do outro grupo."),
        ("hiato_quebra_estrutural.png", "Teste tipo Chow: mudança de patamar/inclinação estatisticamente significativa nos 3 eventos testados (reforma trabalhista, reforma da previdência, recessão 2015-16)."),
    ]),
    ("Aprofundamentos: novas variáveis da PNAD", "Informalidade, jornada, alfabetização, desalento", [
        ("informalidade_carteira_assinada.png", "% com carteira assinada: Branca 72%, Negra 63%, Indígena 52% (2026 T2)."),
        ("renda_por_hora.png", "O hiato sobrevive quase intacto na renda POR HORA — não vem de jornada menor, é remuneração por hora menor mesmo."),
        ("alfabetizacao_60mais.png", "Analfabetismo residual (60+ anos): Branca 93%, Negra 80%, Indígena 72% alfabetizados."),
        ("desalento.png", "Entre quem está fora da força de trabalho, Negra e Indígena desistem de procurar emprego a taxas 2-4x maiores que Branca."),
    ]),
    ("Geração: seguindo a mesma coorte, não a mesma faixa etária", "A PNAD é um corte transversal repetido — 'pessoas de 14-17 anos' em 2012 e 2026 são pessoas diferentes", [
        ("hiato_racial_por_geracao.png", "Cada linha é a MESMA coorte de nascimento envelhecendo — o hiato varia bastante entre gerações e ao longo do tempo dentro de cada uma."),
        ("renda_por_geracao_raca.png", "Renda por geração e raça, cada uma na idade em que está hoje."),
    ]),
    ("Quem está no topo 10%? Negra vs. Branca", "Limiar (P90) calculado DENTRO de cada raça, não um corte único pro Brasil", [
        ("topo10_genero.png", "% de mulheres no topo 10% de cada raça — Branca sempre um pouco à frente, mas a distância vem encolhendo."),
        ("topo10_escolaridade.png", "% com Superior completo no topo 10% — o topo dos negros ficou muito mais qualificado (38%→58%), mas ainda atrás do topo dos brancos (~83%)."),
        ("topo10_faixa_etaria.png", "Composição por faixa etária do topo 10%, média dos últimos 8 trimestres."),
        ("topo10_geracao.png", "Composição por geração do topo 10% — o topo dos negros pende um pouco mais para Millennial/Geração Z; o dos brancos, para Baby Boomer/Geração X."),
    ]),
    ("Decomposição dos 4 quartis de renda — Negra vs. Branca", "Generaliza o topo 10% pra toda a distribuição: Q1 (25% mais pobres) a Q4 (25% mais ricos), limiar DENTRO de cada raça", [
        ("quartis_genero.png", "% de mulheres por quartil — maioria entre as mais pobres, minoria entre as mais ricas, nas duas raças; a distância entre raças é maior no Q4."),
        ("quartis_escolaridade.png", "% com Superior completo cresce em todos os quartis nas duas raças, mas o hiato racial se abre dramaticamente no Q4 (~72% Branca vs. ~40% Negra)."),
        ("quartis_faixa_etaria.png", "Composição por faixa etária, um painel por quartil — perfil etário parecido entre as raças em cada quartil."),
        ("quartis_geracao.png", "Composição por geração, um painel por quartil."),
    ]),
    ("Desigualdade interna, setor público/privado e sobre-qualificação", "Sugestões de análise validadas antes de executar", [
        ("gini_por_raca.png", "Gini DENTRO de cada raça — Branca tem mais desigualdade interna que Negra (distribuição mais comprimida na base, não 'melhor situação')."),
        ("theil_decomposicao.png", "Decomposição de Theil: só 7% da desigualdade total do Brasil vem de diferença ENTRE raças — 93% é desigualdade DENTRO de cada uma. Não diminui o hiato racial, mostra que a desigualdade brasileira tem várias causas."),
        ("hiato_setor_publico_privado.png", "Confirma a hipótese: o hiato racial é menor no setor Público (~45-50%, tabela salarial padronizada) do que no Privado (~58-65%)."),
        ("segregacao_setorial.png", "Segregação por setor econômico (10%) é menor que por ocupação (17%) — as raças se distribuem mais parecido entre setores do que entre cargos dentro deles."),
        ("sobrequalificacao.png", "Negros com Superior completo têm taxa de 'sobre-qualificação' (acabar em ocupação elementar) quase o dobro da de brancos com o mesmo diploma."),
    ]),
    ("A função quantil da renda, em R$, e sua inversa", "'Os 10% mais pobres entre os negros ganham quanto comparado aos 10% mais pobres entre os brancos?'", [
        ("funcao_quantil_racial.png", "Em R$: no P10, Branca ganha R$1.500 e Negra R$720. No P50 (mediana), R$3.000 vs. R$2.000. No P90, R$10.000 vs. R$5.000."),
        ("hiato_por_percentil.png", "O hiato bruto (sem controles) por percentil não é uniforme — é maior na base (108% no P10) e no topo (100% no P90) do que no meio da distribuição."),
        ("percentil_de_valor_racial.png", "A pergunta inversa: quem ganha R$3.000 está no P57 entre os brancos (mediano/meio da distribuição), mas já no P76 entre os negros (quase o topo)."),
    ]),
    ("Raça × Gênero", "Combinado + um por gênero", [
        ("renda_por_raca_genero.png", "O hiato de gênero soma ao de raça, não substitui."),
        ("renda_por_raca_homens.png", "Recorte só Homens."),
        ("renda_por_raca_mulheres.png", "Recorte só Mulheres."),
    ]),
    ("Preta × Parda × Gênero", "Combinado + um por gênero", [
        ("renda_preta_parda_genero.png", "Preta e Parda por gênero — quatro linhas, mesmo padrão de proximidade."),
        ("renda_preta_parda_homens.png", "Recorte só Homens."),
        ("renda_preta_parda_mulheres.png", "Recorte só Mulheres."),
    ]),
    ("Raça × Faixa etária", "Combinado + um por faixa etária", [
        ("renda_por_raca_faixa_etaria.png", "Hiato racial se abre na faixa de maior potencial de renda (40-59 anos)."),
        ("renda_por_raca_faixa_14_17.png", "14-17 anos."),
        ("renda_por_raca_faixa_18_24.png", "18-24 anos."),
        ("renda_por_raca_faixa_25_39.png", "25-39 anos."),
        ("renda_por_raca_faixa_40_59.png", "40-59 anos — maior hiato."),
        ("renda_por_raca_faixa_60mais.png", "60+ anos."),
    ]),
    ("Preta × Parda × Faixa etária", "Combinado + uma por faixa etária", [
        ("renda_preta_parda_faixa_etaria.png", "Preta vs. Parda por faixa etária, snapshot mais recente."),
        ("renda_preta_parda_faixa_14_17.png", "14-17 anos."),
        ("renda_preta_parda_faixa_18_24.png", "18-24 anos."),
        ("renda_preta_parda_faixa_25_39.png", "25-39 anos."),
        ("renda_preta_parda_faixa_40_59.png", "40-59 anos."),
        ("renda_preta_parda_faixa_60mais.png", "60+ anos."),
    ]),
    ("Raça × Escolaridade", "O achado central da Fase 1", [
        ("renda_por_raca_escolaridade.png", "O hiato racial SOBREVIVE ao controle por escolaridade — abre no Superior completo."),
        ("renda_por_raca_escolaridade_sem_instrucao.png", "Sem instrução."),
        ("renda_por_raca_escolaridade_fundamental_incompl.png", "Fundamental incompleto."),
        ("renda_por_raca_escolaridade_fundamental_compl.png", "Fundamental completo."),
        ("renda_por_raca_escolaridade_medio_incompl.png", "Médio incompleto."),
        ("renda_por_raca_escolaridade_medio_compl.png", "Médio completo."),
        ("renda_por_raca_escolaridade_superior_incompl.png", "Superior incompleto."),
        ("renda_por_raca_escolaridade_superior_compl.png", "Superior completo — maior hiato de todos."),
    ]),
    ("Preta × Parda × Escolaridade", "Combinado + um por nível", [
        ("renda_preta_parda_escolaridade.png", "Preta vs. Parda por nível de instrução, snapshot mais recente."),
        ("renda_preta_parda_escolaridade_sem_instrucao.png", "Sem instrução."),
        ("renda_preta_parda_escolaridade_fundamental_incompl.png", "Fundamental incompleto."),
        ("renda_preta_parda_escolaridade_fundamental_compl.png", "Fundamental completo."),
        ("renda_preta_parda_escolaridade_medio_incompl.png", "Médio incompleto."),
        ("renda_preta_parda_escolaridade_medio_compl.png", "Médio completo."),
        ("renda_preta_parda_escolaridade_superior_incompl.png", "Superior incompleto."),
        ("renda_preta_parda_escolaridade_superior_compl.png", "Superior completo."),
    ]),
    ("Raça × Gênero × Escolaridade", None, [
        ("renda_por_raca_escolaridade_homens.png", "Hiato no Superior completo é maior entre homens (~46%)."),
        ("renda_por_raca_escolaridade_mulheres.png", "...do que entre mulheres (~36%)."),
    ]),
    ("Raça × Gênero × Faixa etária", None, [
        ("renda_por_faixa_etaria_raca_genero.png", "Mesmo efeito de abertura em 40-59 anos, visível nos dois gêneros."),
    ]),
    ("Escolaridade × Raça × Gênero", None, [
        ("escolaridade_por_raca_genero.png", "Conclusão do superior: hiato racial >2x; mulheres à frente dos homens em todos os grupos."),
    ]),
    ("Todas as dimensões de uma vez", "Raça × Gênero × Faixa etária × Escolaridade", [
        ("renda_completa_heatmap.png", "Seis painéis (raça×gênero), faixa etária × nível de instrução em cada um — o padrão racial se mantém."),
    ]),
]


def _slide_em_branco(prs: Presentation):
    layout_em_branco = prs.slide_layouts[6]
    return prs.slides.add_slide(layout_em_branco)


def _fundo(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = SUPERFICIE


def slide_titulo(prs: Presentation) -> None:
    slide = _slide_em_branco(prs)
    _fundo(slide)
    caixa = slide.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.7), Inches(2))
    tf = caixa.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Datahub de Análise Racial no Brasil"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = TINTA_PRIMARIA

    p2 = tf.add_paragraph()
    p2.text = "Fase 1 — PNAD Contínua (2012–2026) · Análise exploratória"
    p2.font.size = Pt(20)
    p2.font.color.rgb = TINTA_SECUNDARIA
    p2.space_before = Pt(12)

    p3 = tf.add_paragraph()
    p3.text = "Fonte: IBGE, PNAD Contínua Trimestral (microdados). Elaboração própria."
    p3.font.size = Pt(13)
    p3.font.color.rgb = TINTA_SECUNDARIA
    p3.space_before = Pt(24)


def slide_secao(prs: Presentation, titulo: str, nota: str | None) -> None:
    slide = _slide_em_branco(prs)
    _fundo(slide)
    caixa = slide.shapes.add_textbox(Inches(0.8), Inches(3.1), Inches(11.7), Inches(1.5))
    tf = caixa.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = titulo
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = COR_BRANCA
    if nota:
        p2 = tf.add_paragraph()
        p2.text = nota
        p2.font.size = Pt(16)
        p2.font.color.rgb = TINTA_SECUNDARIA
        p2.space_before = Pt(10)


def slide_grafico(prs: Presentation, arquivo: str, legenda: str) -> None:
    caminho = IMG_DIR / arquivo
    if not caminho.exists():
        print(f"  aviso: {arquivo} não encontrado, pulando slide")
        return
    slide = _slide_em_branco(prs)
    _fundo(slide)

    # imagem ocupa a maior parte do slide, centralizada, com margem pra legenda embaixo
    from PIL import Image
    with Image.open(caminho) as im:
        largura_px, altura_px = im.size
    proporcao = largura_px / altura_px

    max_largura = Inches(12.3)
    max_altura = Inches(6.1)
    if max_largura / proporcao <= max_altura:
        largura = max_largura
        altura = Emu(int(max_largura / proporcao))
    else:
        altura = max_altura
        largura = Emu(int(max_altura * proporcao))

    left = Emu(int((LARGURA_SLIDE - largura) / 2))
    top = Inches(0.35)
    slide.shapes.add_picture(str(caminho), left, top, width=largura, height=altura)

    caixa = slide.shapes.add_textbox(Inches(0.6), Inches(6.65), Inches(12.1), Inches(0.7))
    tf = caixa.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = legenda
    p.font.size = Pt(14)
    p.font.color.rgb = TINTA_SECUNDARIA
    p.alignment = PP_ALIGN.CENTER


def main() -> None:
    prs = Presentation()
    prs.slide_width = LARGURA_SLIDE
    prs.slide_height = ALTURA_SLIDE

    slide_titulo(prs)

    total_graficos = 0
    for titulo_secao, nota, graficos in SECOES:
        slide_secao(prs, titulo_secao, nota)
        for arquivo, legenda in graficos:
            slide_grafico(prs, arquivo, legenda)
            total_graficos += 1

    prs.save(DESTINO)
    total_slides = 1 + len(SECOES) + total_graficos
    print(f"Apresentação salva em: {DESTINO.relative_to(REPO_ROOT)}")
    print(f"Slides: {total_slides} ({total_graficos} gráficos + {len(SECOES)} divisores + 1 título)")


if __name__ == "__main__":
    main()
